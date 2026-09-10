#!/usr/bin/env python3
"""Regenerate all 14 canonical IndPenSim paper figures at 600 DPI.

The script reads the checked-in result/data files from:
https://github.com/abdulbasitbehlim/IndPenSim-penicillin-soft-sensor

Scientific rule: numerical values and evaluation subsets are not recomputed or altered,
except the deterministic baseline Random Forest is refit solely to recover its saved-study
feature-importance vector (the repository does not store that vector as a table).

Outputs:
  paper_figures_600dpi/*.png     (14 publication-style PNGs, 600 DPI)
  IndPenSim_14_Paper_Figures_600dpi.pptx
  IndPenSim_14_Paper_Figures_600dpi.zip
"""
from __future__ import annotations

import json
import math
import os
import shutil
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib as mpl
import matplotlib.pyplot as plt
from matplotlib.colors import TwoSlopeNorm
from matplotlib.ticker import MaxNLocator
from sklearn.ensemble import RandomForestRegressor
from sklearn.impute import SimpleImputer
from sklearn.pipeline import Pipeline
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

warnings.filterwarnings("ignore")

ROOT = Path(__file__).resolve().parent
if not (ROOT / "results").exists():
    # Also support running from scripts/ in a repository checkout.
    candidate = ROOT.parent
    if (candidate / "results").exists():
        ROOT = candidate

OUT = ROOT / "paper_figures_600dpi"
OUT.mkdir(parents=True, exist_ok=True)
DPI = 600

# Colorblind-friendly Okabe-Ito-inspired palette.
BLUE = "#0072B2"
ORANGE = "#D55E00"
GREEN = "#009E73"
PURPLE = "#CC79A7"
SKY = "#56B4E9"
YELLOW = "#F0E442"
BLACK = "#111111"
DARK_GREY = "#555555"
MID_GREY = "#9A9A9A"
LIGHT_GREY = "#D9D9D9"
GRID_GREY = "#E9E9E9"

mpl.rcParams.update({
    "figure.dpi": 160,
    "savefig.dpi": DPI,
    "font.family": "DejaVu Sans",
    "font.size": 9.5,
    "axes.labelsize": 10.5,
    "axes.titlesize": 11,
    "xtick.labelsize": 9,
    "ytick.labelsize": 9,
    "legend.fontsize": 8.5,
    "axes.linewidth": 0.8,
    "lines.linewidth": 1.8,
    "pdf.fonttype": 42,
    "ps.fonttype": 42,
})

TARGET = "Penicillin concentration(P:g/L)"
TIME = "Time (h)"
BATCH = "Batch_ID"


def read_csv(path: str) -> pd.DataFrame:
    return pd.read_csv(ROOT / path)


def save(fig: plt.Figure, name: str) -> Path:
    path = OUT / name
    fig.savefig(path, dpi=DPI, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def clean_ax(ax: plt.Axes, grid=True) -> None:
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color(DARK_GREY)
    ax.spines["bottom"].set_color(DARK_GREY)
    ax.tick_params(colors=DARK_GREY, width=0.8)
    if grid:
        ax.grid(axis="y", color=GRID_GREY, linewidth=0.7)
        ax.set_axisbelow(True)


def annotate_bars(ax, bars, fmt="{:.2f}", dy=0.02, fontsize=8.5):
    ylim = ax.get_ylim()
    span = ylim[1] - ylim[0]
    for b in bars:
        y = b.get_height()
        va = "bottom" if y >= 0 else "top"
        offset = dy * span if y >= 0 else -dy * span
        ax.text(b.get_x() + b.get_width()/2, y + offset, fmt.format(y),
                ha="center", va=va, fontsize=fontsize, color=DARK_GREY)


def figure_b1() -> Path:
    train = read_csv("data/splits/train_normal_60_batches.csv")
    batches = [30, 17, 8]
    colors = [BLUE, ORANGE, GREEN]
    fig, ax = plt.subplots(figsize=(6.6, 4.5))
    for bid, c in zip(batches, colors):
        d = train[train[BATCH] == bid].sort_values(TIME)
        ax.plot(d[TIME], d[TARGET], color=c, label=f"Batch {bid}")
    ax.set_xlabel("Time (h)")
    ax.set_ylabel("Penicillin concentration (g/L)")
    ax.legend(frameon=False, ncol=3, loc="upper left")
    clean_ax(ax)
    return save(fig, "B1_training_batch_trajectories_600dpi.png")


def figure_b2() -> Path:
    normal = read_csv("results/baseline/normal_test_predictions.csv")
    fault = read_csv("results/baseline/fault_test_predictions.csv")
    fig, axes = plt.subplots(1, 2, figsize=(9.0, 4.35), sharex=True, sharey=True)
    limits = [0, max(normal[TARGET].max(), fault[TARGET].max(), normal["Predicted_Target"].max(), fault["Predicted_Target"].max()) * 1.03]
    for ax, data, color, title, letter in [
        (axes[0], normal, BLUE, "Normal test", "a"),
        (axes[1], fault, ORANGE, "Fault stress test", "b"),
    ]:
        ax.scatter(data[TARGET], data["Predicted_Target"], s=3.2, alpha=0.23, color=color, linewidths=0)
        ax.plot(limits, limits, "--", color=BLACK, linewidth=1.1, label="Perfect prediction")
        ax.set_xlim(limits); ax.set_ylim(limits)
        ax.set_xlabel("Actual concentration (g/L)")
        ax.set_ylabel("Predicted concentration (g/L)")
        ax.text(0.04, 0.94, title, transform=ax.transAxes, ha="left", va="top", weight="bold")
        ax.text(-0.12, 1.04, letter, transform=ax.transAxes, weight="bold", fontsize=12)
        clean_ax(ax)
    axes[0].legend(frameon=False, loc="lower right")
    fig.tight_layout(w_pad=2.0)
    return save(fig, "B2_baseline_actual_vs_predicted_600dpi.png")


def add_history_features(frame: pd.DataFrame) -> pd.DataFrame:
    df = frame.copy()
    history_cols = [
        "Dissolved oxygen concentration(DO2:mg/L)",
        "Sugar feed rate(Fs:L/h)",
        "Temperature(T:K)",
        "pH(pH:pH)",
        "Aeration rate(Fg:L/h)",
    ]
    for col in history_cols:
        g = df.groupby(BATCH, sort=False)[col]
        df[col + "_lag1"] = g.shift(1)
        df[col + "_difference1"] = g.diff(1)
        df[col + "_previous5_mean"] = g.transform(lambda s: s.shift(1).rolling(window=5, min_periods=1).mean())
    dt = df.groupby(BATCH, sort=False)[TIME].diff()
    # Notebook behavior: the first row adds zero feed, subsequent rows use elapsed time.
    dt = dt.fillna(0).clip(lower=0)
    df["Cumulative_Sugar_Feed"] = (df["Sugar feed rate(Fs:L/h)"] * dt).groupby(df[BATCH]).cumsum()
    return df


def figure_b3() -> Path:
    info = json.loads((ROOT / "results/baseline/model_information.json").read_text())
    features = info["model_features"]
    train = read_csv("data/splits/train_normal_60_batches.csv")
    val = read_csv("data/splits/validation_normal_15_batches.csv")
    development = pd.concat([train, val], ignore_index=True)
    development = add_history_features(development)
    X = development[features]
    y = development[TARGET]
    model = Pipeline([
        ("imputer", SimpleImputer(strategy="median")),
        ("model", RandomForestRegressor(
            n_estimators=150, max_depth=18, min_samples_leaf=2,
            random_state=42, n_jobs=-1,
        )),
    ])
    model.fit(X, y)
    imp = pd.Series(model.named_steps["model"].feature_importances_, index=features).sort_values(ascending=False).head(15)

    def friendly(s: str) -> str:
        repl = {
            "Cumulative_Sugar_Feed": "Cumulative feed volume",
            "Time (h)": "Time",
            "Oxygen in percent in off-gas(O2:O2  (%))": "Off-gas O₂",
            "carbon dioxide percent in off-gas(CO2outgas:%)": "Off-gas CO₂",
            "Dissolved oxygen concentration(DO2:mg/L)": "Dissolved oxygen",
            "Vessel Volume(V:L)": "Vessel volume",
            "Vessel Weight(Wt:Kg)": "Vessel weight",
            "Oxygen Uptake Rate(OUR:(g min^{-1}))": "Oxygen uptake rate",
        }
        out = repl.get(s, s)
        out = out.replace("Dissolved oxygen concentration(DO2:mg/L)", "Dissolved oxygen")
        out = out.replace("Sugar feed rate(Fs:L/h)", "Sugar feed rate")
        out = out.replace("_previous5_mean", " (previous-5 mean)").replace("_lag1", " (lag 1)").replace("_difference1", " (difference 1)")
        return out

    labels = [friendly(x) for x in imp.index]
    fig, ax = plt.subplots(figsize=(8.3, 5.05))
    vals = imp.values[::-1]
    labels = labels[::-1]
    colors = [SKY] * len(vals)
    if len(colors) >= 2:
        colors[-1] = BLUE; colors[-2] = BLUE
    ax.barh(range(len(vals)), vals, color=colors, edgecolor="none")
    ax.set_yticks(range(len(vals)), labels)
    ax.set_xlabel("Random Forest impurity-based importance")
    clean_ax(ax, grid=False)
    ax.grid(axis="x", color=GRID_GREY, linewidth=0.7); ax.set_axisbelow(True)
    return save(fig, "B3_random_forest_feature_importance_600dpi.png")


def figure_e1() -> Path:
    df = read_csv("results/experiments/experiment_1_repeated_batch_cv_fold_metrics.csv")
    fig, axes = plt.subplots(1, 2, figsize=(8.8, 4.45))
    rng = np.random.default_rng(42)
    for ax, col, ylabel, color, letter in [
        (axes[0], "Row_RMSE", "Pooled RMSE (g/L)", BLUE, "a"),
        (axes[1], "Row_R2", "Pooled R²", GREEN, "b"),
    ]:
        vals = df[col].astype(float).to_numpy()
        bp = ax.boxplot(vals, positions=[1], widths=0.46, patch_artist=True,
                        boxprops=dict(facecolor=mpl.colors.to_rgba(color, 0.28), edgecolor=color, linewidth=1.1),
                        medianprops=dict(color=BLACK, linewidth=1.5),
                        whiskerprops=dict(color=DARK_GREY), capprops=dict(color=DARK_GREY),
                        flierprops=dict(marker="o", markersize=3, markerfacecolor=color, markeredgecolor="none", alpha=0.7))
        x = 1 + rng.normal(0, 0.025, size=len(vals))
        ax.scatter(x, vals, s=17, color=color, alpha=0.72, zorder=3, linewidths=0)
        ax.set_xticks([1], ["25 evaluations"])
        ax.set_ylabel(ylabel)
        ax.text(-0.1, 1.04, letter, transform=ax.transAxes, weight="bold", fontsize=12)
        clean_ax(ax)
    fig.tight_layout(w_pad=2.5)
    return save(fig, "E1_repeated_batch_cv_distribution_600dpi.png")


def figure_e2() -> Path:
    df = read_csv("results/experiments/experiment_2_time_feed_ablation.csv")
    variants = list(dict.fromkeys(df["Variant"].tolist()))
    normal = [float(df[(df["Variant"] == v) & (df["Dataset"].str.contains("Normal", case=False))]["RMSE"].iloc[0]) for v in variants]
    fault = [float(df[(df["Variant"] == v) & (df["Dataset"].str.contains("Fault", case=False))]["RMSE"].iloc[0]) for v in variants]
    labels = ["All features", "− Time", "− cumulative feed", "− Time and cumulative feed"]
    x = np.arange(len(variants)); width = 0.32
    fig, ax = plt.subplots(figsize=(7.7, 4.6))
    b1 = ax.bar(x-width/2, normal, width, label="Normal test", color=BLUE)
    b2 = ax.bar(x+width/2, fault, width, label="Fault test", color=ORANGE)
    ax.set_ylabel("RMSE (g/L)")
    ax.set_xticks(x, labels, rotation=20, ha="right")
    ax.legend(frameon=False, ncol=2)
    clean_ax(ax)
    ax.set_ylim(0, max(fault)*1.15)
    annotate_bars(ax, b1); annotate_bars(ax, b2)
    return save(fig, "E2_time_feed_ablation_600dpi.png")


def figure_e3() -> Path:
    byb = read_csv("results/experiments/experiment_3_fault_phase_by_batch.csv")
    overall = read_csv("results/experiments/experiment_3_fault_phase_overall.csv")
    phases = ["Before fault", "During fault window", "After fault"]
    short = ["Before", "During", "After"]
    x = np.arange(3)
    fig, ax = plt.subplots(figsize=(6.4, 4.55))
    pivot = byb.pivot(index="Batch_ID", columns="Fault_Phase", values="MAE")
    for _, row in pivot.iterrows():
        ys = [row.get(p, np.nan) for p in phases]
        ax.plot(x, ys, color=LIGHT_GREY, linewidth=1.0, marker="o", markersize=2.8, alpha=0.85, zorder=1)
    lookup = overall.set_index("Fault_Phase")
    means = np.array([float(lookup.loc[p, "Mean_Batch_MAE"]) for p in phases])
    lo = np.array([float(lookup.loc[p, "Bootstrap_95pct_Lower_Mean_Batch_MAE"]) for p in phases])
    hi = np.array([float(lookup.loc[p, "Bootstrap_95pct_Upper_Mean_Batch_MAE"]) for p in phases])
    yerr = np.vstack([means-lo, hi-means])
    ax.errorbar(x, means, yerr=yerr, color=ORANGE, marker="o", markersize=5.5,
                capsize=4, linewidth=2.0, label="Equal-batch mean ± 95% bootstrap CI", zorder=4)
    for xx, yy in zip(x, means):
        ax.text(xx, yy + 0.35, f"{yy:.3f}", ha="center", color=DARK_GREY, fontsize=8.5)
    ax.set_xticks(x, short)
    ax.set_ylabel("Batch MAE (g/L)")
    ax.legend(frameon=False, loc="upper left")
    clean_ax(ax)
    return save(fig, "E3_fault_phase_mae_600dpi.png")


def figure_e4() -> Path:
    df = read_csv("results/experiments/experiment_4_model_comparison_test.csv")
    models = ["Dummy Median", "Linear Regression", "Random Forest", "HistGradientBoosting"]
    fig, axes = plt.subplots(1, 2, figsize=(9.7, 4.25))
    for ax, metric, ylabel, letter in [(axes[0], "RMSE", "Pooled RMSE (g/L)", "a"), (axes[1], "R2", "Pooled R²", "b")]:
        n = [float(df[(df["Model"]==m) & (df["Dataset"].str.contains("Normal", case=False))][metric].iloc[0]) for m in models]
        f = [float(df[(df["Model"]==m) & (df["Dataset"].str.contains("Fault", case=False))][metric].iloc[0]) for m in models]
        x=np.arange(len(models)); w=0.34
        b1=ax.bar(x-w/2,n,w,color=BLUE,label="Normal test")
        b2=ax.bar(x+w/2,f,w,color=ORANGE,label="Fault test")
        ax.set_xticks(x, models, rotation=22, ha="right")
        ax.set_ylabel(ylabel)
        ax.text(-0.1, 1.04, letter, transform=ax.transAxes, weight="bold", fontsize=12)
        clean_ax(ax)
        if metric == "RMSE":
            ax.set_ylim(0, max(f)*1.16); annotate_bars(ax,b1); annotate_bars(ax,b2)
        else:
            ax.axhline(0,color=DARK_GREY,linewidth=0.8)
            ymin=min(min(n),min(f)); ymax=max(max(n),max(f)); ax.set_ylim(ymin-0.08,ymax+0.12)
            annotate_bars(ax,b1,fmt="{:.2f}",dy=0.015); annotate_bars(ax,b2,fmt="{:.2f}",dy=0.015)
        ax.legend(frameon=False, ncol=2, loc="best")
    fig.tight_layout(w_pad=2.2)
    return save(fig, "E4_model_comparison_600dpi.png")


def figure_e5a() -> Path:
    df = read_csv("results/experiments/experiment_5_early_ood_scores.csv")
    sub = df[(df["Set"] == "Fault test") & (df["Horizon_h"].isin([12,24,48,72]))].copy()
    pivot = sub.pivot(index="Batch_ID", columns="Horizon_h", values="OOD_Excess_Above_Threshold").reindex(index=range(91,101), columns=[12,24,48,72])
    data = pivot.to_numpy(float)
    lim = max(abs(np.nanmin(data)), abs(np.nanmax(data)))
    fig, ax = plt.subplots(figsize=(6.1, 4.65))
    im = ax.imshow(data, aspect="auto", cmap="RdBu_r", norm=TwoSlopeNorm(vcenter=0, vmin=-lim, vmax=lim))
    ax.set_xticks(range(4), ["12 h","24 h","48 h","72 h"])
    ax.set_yticks(range(10), [str(i) for i in range(91,101)])
    ax.set_xlabel("Observed batch duration")
    ax.set_ylabel("Fault batch ID")
    for i in range(data.shape[0]):
        for j in range(data.shape[1]):
            v=data[i,j]
            rgba=im.cmap(im.norm(v)); lum=0.299*rgba[0]+0.587*rgba[1]+0.114*rgba[2]
            ax.text(j,i,f"{v:+.3f}",ha="center",va="center",fontsize=7.6,color="white" if lum<0.5 else DARK_GREY)
    cbar=fig.colorbar(im,ax=ax,pad=0.04)
    cbar.set_label("OOD score − training threshold")
    for s in ax.spines.values(): s.set_visible(False)
    return save(fig, "E5a_early_ood_heatmap_600dpi.png")


def figure_e5b() -> Path:
    df = read_csv("results/experiments/experiment_5_ood_uncertainty_by_batch.csv")
    # One record per evaluated batch at the 24-h OOD horizon.
    if "Horizon_h" in df.columns:
        df = df[df["Horizon_h"] == 24].copy()
    elif df[BATCH].duplicated().any():
        df = df.drop_duplicates(BATCH, keep="first")
    fig, ax = plt.subplots(figsize=(6.45,4.55))
    for label, color, marker in [("Normal test",BLUE,"o"),("Fault test",ORANGE,"s")]:
        d=df[df["Set"]==label]
        ax.scatter(d["OOD_Score"],d["RMSE"],s=34,color=color,marker=marker,label=label,alpha=0.88,edgecolor="white",linewidth=0.4)
    threshold=float(df["Training_95pct_Threshold"].dropna().iloc[0])
    ax.axvline(threshold,color=BLACK,linestyle="--",linewidth=1.1,label="Training 95th-percentile threshold")
    for bid in [91,100]:
        r=df[df[BATCH]==bid].iloc[0]
        ax.annotate(f"Batch {bid}",(r["OOD_Score"],r["RMSE"]),xytext=(5,6),textcoords="offset points",fontsize=8,color=DARK_GREY)
    ax.set_xlabel("24 h OOD score")
    ax.set_ylabel("Full-batch RMSE (g/L)")
    ax.legend(frameon=False,loc="upper left")
    clean_ax(ax)
    return save(fig, "E5b_ood_vs_batch_rmse_600dpi.png")


def normalize_model_names(s: pd.Series) -> pd.Series:
    return s.replace({"Fault-aware HGB":"Fault-inclusive HGB"})


def figure_f1() -> Path:
    df=read_csv("results/main/cross_validated_overall_metrics.csv")
    df=df.copy(); df["DisplayModel"]=normalize_model_names(df["Model"])
    conditions=["Normal","Fault"]
    labels=["Held-out normal","Held-out fault"]
    models=["Normal-only HGB","Fault-inclusive HGB"]
    vals={m:[float(df[(df["Condition"]==c)&(df["DisplayModel"]==m)]["RMSE"].iloc[0]) for c in conditions] for m in models}
    x=np.arange(2); w=0.32
    fig,ax=plt.subplots(figsize=(6.0,4.0))
    b1=ax.bar(x-w/2,vals[models[0]],w,color=BLUE,label=models[0])
    b2=ax.bar(x+w/2,vals[models[1]],w,color=ORANGE,label=models[1])
    ax.set_xticks(x,labels); ax.set_ylabel("RMSE (g/L)")
    ax.legend(frameon=False,ncol=2,loc="upper left")
    ax.set_ylim(0,max(max(vals[models[0]]),max(vals[models[1]]))*1.2)
    clean_ax(ax); annotate_bars(ax,b1,fmt="{:.3f}"); annotate_bars(ax,b2,fmt="{:.3f}")
    return save(fig,"F1_overall_rmse_comparison_600dpi.png")


def figure_f2() -> Path:
    df=read_csv("results/main/cross_validated_batch_metrics.csv")
    df=df[df["Condition"]=="Fault"].copy(); df["DisplayModel"]=normalize_model_names(df["Model"])
    bids=list(range(91,101)); x=np.arange(len(bids)); w=0.36
    n=[float(df[(df[BATCH]==b)&(df["DisplayModel"]=="Normal-only HGB")]["RMSE"].iloc[0]) for b in bids]
    f=[float(df[(df[BATCH]==b)&(df["DisplayModel"]=="Fault-inclusive HGB")]["RMSE"].iloc[0]) for b in bids]
    fig,ax=plt.subplots(figsize=(7.25,4.35))
    ax.bar(x-w/2,n,w,color=BLUE,label="Normal-only HGB")
    ax.bar(x+w/2,f,w,color=ORANGE,label="Fault-inclusive HGB")
    ax.set_xticks(x,[str(b) for b in bids])
    ax.set_xlabel("Batch ID"); ax.set_ylabel("Batch RMSE (g/L)")
    ax.legend(frameon=False,ncol=2,loc="upper left")
    clean_ax(ax)
    # Deliberately no plot title: the manuscript caption carries the figure title.
    return save(fig,"F2_fault_batch_rmse_600dpi.png")


def figure_f3() -> Path:
    df=read_csv("results/main/cross_validated_predictions.csv")
    d=df[df["Condition"]=="Fault"].copy()
    fig,axes=plt.subplots(1,2,figsize=(9.0,4.35),sharex=True,sharey=True)
    maxv=max(d[TARGET].max(),d["Normal_Only_Prediction"].max(),d["Fault_Aware_Prediction"].max())*1.03
    limits=[0,maxv]
    for ax,col,color,title,letter in [
        (axes[0],"Normal_Only_Prediction",BLUE,"Normal-only HGB","a"),
        (axes[1],"Fault_Aware_Prediction",ORANGE,"Fault-inclusive HGB","b"),
    ]:
        ax.scatter(d[TARGET],d[col],s=3.2,alpha=0.22,color=color,linewidths=0)
        ax.plot(limits,limits,"--",color=BLACK,linewidth=1.1,label="Perfect prediction")
        ax.set_xlim(limits); ax.set_ylim(limits)
        ax.set_xlabel("Actual concentration (g/L)"); ax.set_ylabel("Predicted concentration (g/L)")
        ax.text(0.04,0.94,title,transform=ax.transAxes,ha="left",va="top",weight="bold")
        ax.text(-0.12,1.04,letter,transform=ax.transAxes,weight="bold",fontsize=12)
        clean_ax(ax)
    axes[0].legend(frameon=False,loc="lower right")
    fig.tight_layout(w_pad=2)
    return save(fig,"F3_fault_actual_vs_predicted_600dpi.png")


def figure_f4() -> Path:
    df=read_csv("results/main/cross_validated_predictions.csv")
    normal=df[df["Condition"]=="Normal"]["Fault_Risk_Probability"].astype(float).to_numpy()
    fault=df[df["Condition"]=="Fault"].copy()
    pre=fault[fault["Fault_Affected"].astype(int)==0]["Fault_Risk_Probability"].astype(float).to_numpy()
    affected=fault[fault["Fault_Affected"].astype(int)==1]["Fault_Risk_Probability"].astype(float).to_numpy()
    data=[normal,pre,affected]; labels=["Normal","Fault batch\npre-onset","Fault batch\nonset/after"]
    fig,ax=plt.subplots(figsize=(6.0,4.15))
    bp=ax.boxplot(data,labels=labels,patch_artist=True,widths=0.55,showfliers=False,
                  medianprops=dict(color=BLACK,linewidth=1.4),whiskerprops=dict(color=DARK_GREY),capprops=dict(color=DARK_GREY))
    for patch,c in zip(bp["boxes"],[BLUE,SKY,ORANGE]):
        patch.set_facecolor(mpl.colors.to_rgba(c,0.35)); patch.set_edgecolor(c); patch.set_linewidth(1.0)
    ax.axhline(0.5,color=BLACK,linestyle="--",linewidth=1.0,label="Warning threshold = 0.5")
    ax.set_ylabel("Fault-risk warning score (uncalibrated)")
    ax.set_ylim(-0.03,1.03); ax.legend(frameon=False,loc="upper left")
    clean_ax(ax)
    return save(fig,"F4_fault_risk_by_phase_600dpi.png")


def figure_f5() -> Path:
    df=read_csv("results/main/cross_validated_predictions.csv")
    d=df[(df["Condition"]=="Fault")&(df[BATCH]==100)].sort_values(TIME).copy()
    fig,ax=plt.subplots(figsize=(7.1,4.15))
    ax.plot(d[TIME],d[TARGET],color=BLACK,label="Actual",linewidth=2.0)
    ax.plot(d[TIME],d["Normal_Only_Prediction"],color=BLUE,label="Normal-only HGB",linewidth=1.6)
    ax.plot(d[TIME],d["Fault_Aware_Prediction"],color=ORANGE,label="Fault-inclusive HGB",linewidth=1.6)
    affected=d[d["Fault_Affected"].astype(int)==1]
    if not affected.empty:
        onset=float(affected[TIME].iloc[0])
        ax.axvline(onset,color=PURPLE,linestyle="--",linewidth=1.2,label=f"Recorded fault onset ({onset:g} h)")
    ax.set_xlabel("Time (h)"); ax.set_ylabel("Penicillin concentration (g/L)")
    ax.legend(frameon=False,ncol=2,loc="best")
    clean_ax(ax)
    return save(fig,"F5_batch_100_held_out_trajectory_600dpi.png")


FIGURES = [
    ("B1", "Training-batch concentration trajectories", figure_b1),
    ("B2", "Baseline Random Forest: actual versus predicted", figure_b2),
    ("B3", "Baseline Random Forest feature importance", figure_b3),
    ("E1", "Repeated complete-batch cross-validation", figure_e1),
    ("E2", "Time and cumulative-feed ablation", figure_e2),
    ("E3", "Error before, during and after the fault window", figure_e3),
    ("E4", "Stronger model comparison", figure_e4),
    ("E5a", "Early out-of-distribution warnings", figure_e5a),
    ("E5b", "Baseline RF: 24 h OOD score vs full-batch RMSE", figure_e5b),
    ("F1", "Overall held-out RMSE", figure_f1),
    ("F2", "RMSE for each held-out fault batch", figure_f2),
    ("F3", "Fault-batch actual versus predicted concentration", figure_f3),
    ("F4", "Fault-risk warning score by phase", figure_f4),
    ("F5", "Batch 100 held-out trajectory", figure_f5),
]

LEGENDS = {
    "B1": "Representative normal training-batch penicillin concentration trajectories.",
    "B2": "Random Forest baseline predictions on normal and fault test batches; the dashed diagonal denotes perfect agreement.",
    "B3": "Random Forest impurity-based feature importance for the highest-ranked baseline predictors.",
    "E1": "Distribution of pooled RMSE and R² across 25 repeated complete-batch Random Forest evaluations.",
    "E2": "Normal and fault RMSE for all features and the time/cumulative-feed ablation variants.",
    "E3": "Equal-batch mean MAE before, during and after the fault window with the saved bootstrap 95% interval and individual batch traces.",
    "E4": "Fixed-test comparison of Dummy Median, Linear Regression, Random Forest and HistGradientBoosting.",
    "E5a": "OOD score minus the horizon-specific training threshold at 12, 24, 48 and 72 h; positive values are threshold exceedances.",
    "E5b": "Baseline Random Forest/OOD experiment: 24 h raw OOD score versus full-batch RMSE; dashed line is the training 95th-percentile threshold.",
    "F1": "Pooled held-out RMSE for normal-only and fault-inclusive HGB on normal and fault conditions.",
    "F2": "Complete-batch RMSE for held-out fault batches 91–100, comparing normal-only and fault-inclusive HGB.",
    "F3": "Held-out fault-batch actual versus predicted penicillin concentration for the two HGB training strategies.",
    "F4": "Out-of-fold fault-risk warning score by process phase. Individual outlier markers are suppressed for visual clarity; box/median/whisker summaries are unchanged. The score is uncalibrated.",
    "F5": "Held-out batch 100 prediction trajectory; dashed vertical line marks the recorded fault onset.",
}


def add_text(slide, text, x, y, w, h, size, bold=False, color=(35,35,35), align=PP_ALIGN.LEFT):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.clear(); p=tf.paragraphs[0]; p.text=text; p.alignment=align
    r=p.runs[0]; r.font.name="Aptos"; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=RGBColor(*color)
    return box


def fit_picture(slide, image_path: Path, x, y, max_w, max_h):
    from PIL import Image
    with Image.open(image_path) as im:
        aspect=im.width/im.height
    if max_w/max_h < aspect:
        w=max_w; h=w/aspect
    else:
        h=max_h; w=h*aspect
    xx=x+(max_w-w)/2; yy=y+(max_h-h)/2
    slide.shapes.add_picture(str(image_path),Inches(xx),Inches(yy),width=Inches(w),height=Inches(h))


def make_ppt(generated: list[tuple[str,str,Path]]) -> Path:
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    cover=prs.slides.add_slide(prs.slide_layouts[6]); cover.background.fill.solid(); cover.background.fill.fore_color.rgb=RGBColor(255,255,255)
    add_text(cover,"IndPenSim Penicillin Soft Sensor Images",0.8,2.3,11.7,0.9,28,bold=True,color=(25,35,45),align=PP_ALIGN.CENTER)
    add_text(cover,"Publication-ready 600-DPI figure set",0.8,3.2,11.7,0.45,16,color=(0,114,178),align=PP_ALIGN.CENTER)
    for fid,title,path in generated:
        slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor(255,255,255)
        add_text(slide,f"{fid} -- {title}",0.6,0.22,12.1,0.65,21,bold=True,color=(25,35,45))
        fit_picture(slide,path,0.65,1.0,12.0,5.15)
        add_text(slide,"Figure legend. "+LEGENDS[fid],0.65,6.25,12.0,0.75,9.0,color=(55,55,55))
    out=ROOT/"IndPenSim_14_Paper_Figures_600dpi.pptx"; prs.save(out); return out


def main():
    generated=[]
    for fid,title,fn in FIGURES:
        print(f"Generating {fid}: {title}")
        generated.append((fid,title,fn()))
    ppt=make_ppt(generated)
    zip_base=ROOT/"IndPenSim_14_Paper_Figures_600dpi"
    if (ROOT/(zip_base.name+".zip")).exists(): (ROOT/(zip_base.name+".zip")).unlink()
    shutil.make_archive(str(zip_base),"zip",root_dir=OUT)
    print("Generated",len(generated),"600-DPI PNGs")
    print("PowerPoint:",ppt)
    print("Figures:",OUT)

if __name__=="__main__":
    main()
