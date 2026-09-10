#!/usr/bin/env python3
"""Build organized 300-DPI, original-style 600-DPI and publishable 600-DPI figure sets.

Original result PNGs under results/ are never deleted or overwritten.
"""
from pathlib import Path
import csv, hashlib, shutil, subprocess, sys
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT=Path(__file__).resolve().parents[1]
FROOT=ROOT/'figures'; D300=FROOT/'images_300_dpi'; D600=FROOT/'images_600_dpi_original_code_output'; DPUB=FROOT/'publishable_images'
ITEMS=[
('B1','Training-batch concentration trajectories','results/baseline/figures/training_batch_trajectories.png','B1_training_batch_trajectories.png','B1_training_batch_trajectories_600dpi.png'),
('B2','Baseline Random Forest: actual versus predicted','results/baseline/figures/actual_vs_predicted.png','B2_baseline_actual_vs_predicted.png','B2_baseline_actual_vs_predicted_600dpi.png'),
('B3','Baseline Random Forest feature importance','results/baseline/figures/feature_importance.png','B3_random_forest_feature_importance.png','B3_random_forest_feature_importance_600dpi.png'),
('E1','Repeated complete-batch cross-validation','results/experiments/figure_experiment_1_cv_distribution.png','E1_repeated_batch_cv_distribution.png','E1_repeated_batch_cv_distribution_600dpi.png'),
('E2','Time and cumulative-feed ablation','results/experiments/figure_experiment_2_time_feed_ablation.png','E2_time_feed_ablation.png','E2_time_feed_ablation_600dpi.png'),
('E3','Error before, during and after the fault window','results/experiments/figure_experiment_3_fault_phase_mae.png','E3_fault_phase_mae.png','E3_fault_phase_mae_600dpi.png'),
('E4','Stronger model comparison','results/experiments/figure_experiment_4_model_comparison.png','E4_model_comparison.png','E4_model_comparison_600dpi.png'),
('E5a','Early out-of-distribution warnings','results/experiments/figure_experiment_5_fault_ood_heatmap.png','E5a_early_ood_heatmap.png','E5a_early_ood_heatmap_600dpi.png'),
('E5b','Baseline RF: 24 h OOD score vs full-batch RMSE','results/experiments/figure_experiment_5_ood_vs_batch_rmse.png','E5b_ood_vs_batch_rmse.png','E5b_ood_vs_batch_rmse_600dpi.png'),
('F1','Overall held-out RMSE','results/main/figure_1_rmse_comparison.png','F1_overall_rmse_comparison.png','F1_overall_rmse_comparison_600dpi.png'),
('F2','RMSE for each held-out fault batch','results/main/figure_2_fault_batch_rmse.png','F2_fault_batch_rmse.png','F2_fault_batch_rmse_600dpi.png'),
('F3','Fault-batch actual versus predicted concentration','results/main/figure_3_fault_actual_vs_predicted.png','F3_fault_actual_vs_predicted.png','F3_fault_actual_vs_predicted_600dpi.png'),
('F4','Fault-risk warning score by phase','results/main/figure_4_fault_risk_by_phase.png','F4_fault_risk_by_phase.png','F4_fault_risk_by_phase_600dpi.png'),
('F5','Batch 100 held-out trajectory','results/main/batch_100_held_out_trajectory.png','F5_batch_100_held_out_trajectory.png','F5_batch_100_held_out_trajectory_600dpi.png')]

def prep():
    for d in (D300,D600,DPUB):
        d.mkdir(parents=True,exist_ok=True)
        for p in d.glob('*'): p.unlink()

def resave(src,dst,dpi,width=None):
    with Image.open(src) as im:
        out=im.copy()
        if width and out.width<width:
            out=out.resize((width,round(out.height*width/out.width)),Image.Resampling.LANCZOS)
        out.save(dst,'PNG',dpi=(dpi,dpi))

def legacy_sets():
    for _,_,rel,name,_ in ITEMS:
        src=ROOT/rel
        resave(src,D300/name,300)
        resave(src,D600/name,600,4200)

def publication_set():
    gen=FROOT/'paper_figures_600dpi.py'
    code=gen.read_text()
    code=code.replace('info["model_features"]','info["features"]')
    code=code.replace('Bootstrap_95pct_Lower_Mean_Batch_MAE','Mean_Batch_MAE_Bootstrap_95_Lower')
    code=code.replace('Bootstrap_95pct_Upper_Mean_Batch_MAE','Mean_Batch_MAE_Bootstrap_95_Upper')
    gen.write_text(code)
    subprocess.run([sys.executable,str(gen)],cwd=ROOT,check=True)
    out=ROOT/'paper_figures_600dpi'
    for _,_,_,name,gname in ITEMS: shutil.copy2(out/gname,DPUB/name)
    gp=ROOT/'IndPenSim_14_Paper_Figures_600dpi.pptx'
    if gp.exists(): shutil.copy2(gp,DPUB/'Publication_Ready_Figures_600dpi_with_legends.pptx')

def add_text(slide,text,x,y,w,h,size,bold=False):
    b=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); p=b.text_frame.paragraphs[0]; p.text=text
    r=p.runs[0]; r.font.name='Aptos'; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=RGBColor(35,35,35)

def add_pic(slide,path,x,y,mw,mh):
    with Image.open(path) as im: a=im.width/im.height
    if mw/mh<a: w=mw; h=w/a
    else: h=mh; w=h*a
    slide.shapes.add_picture(str(path),Inches(x+(mw-w)/2),Inches(y+(mh-h)/2),width=Inches(w),height=Inches(h))

def deck(folder,path,cover=False):
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    if cover:
        s=prs.slides.add_slide(prs.slide_layouts[6]); add_text(s,'IndPenSim Penicillin Soft Sensor Images',1,2.4,11.3,0.8,28,True); add_text(s,'Publication-ready 600-DPI figure set',1,3.3,11.3,0.5,16)
    for fid,title,_,name,_ in ITEMS:
        s=prs.slides.add_slide(prs.slide_layouts[6]); add_text(s,f'{fid} -- {title}',0.6,0.2,12,0.7,21,True); add_pic(s,folder/name,0.6,1.0,12.1,5.7)
    prs.save(path)

def manifest():
    with (FROOT/'image_manifest.csv').open('w',newline='') as f:
        w=csv.writer(f); w.writerow(['set','figure_id','filename','width_px','height_px','dpi_x','dpi_y','sha256'])
        for setname,folder in [('images_300_dpi',D300),('images_600_dpi_original_code_output',D600),('publishable_images',DPUB)]:
            for fid,_,_,name,_ in ITEMS:
                p=folder/name
                with Image.open(p) as im: dx,dy=im.info.get('dpi',(0,0)); width,height=im.size
                w.writerow([setname,fid,name,width,height,f'{dx:.3f}',f'{dy:.3f}',hashlib.sha256(p.read_bytes()).hexdigest()])

def docs():
    (FROOT/'README.md').write_text('''# High-resolution figure archive\n\nThe original scientific output PNGs under `results/` are preserved unchanged. This archive adds organized copies and manuscript-ready versions.\n\n| Folder | Contents |\n|---|---|\n| `images_300_dpi/` | Organized 300-DPI copies of the previous/canonical code-output figures. |\n| `images_600_dpi_original_code_output/` | 600-DPI archival versions preserving the original output appearance, plus a PowerPoint. |\n| `publishable_images/` | Publication-ready 600-DPI figures generated from saved repository data/results, plus PowerPoint decks. |\n\n`paper_figures_600dpi.py` regenerates the publication figures from checked-in results. `image_manifest.csv` records dimensions, DPI and SHA-256 hashes.\n\nScientific presentation notes: E5a uses three decimals; E5b is explicitly identified as the baseline RF/OOD experiment; F4 is labelled as an uncalibrated warning score and suppresses individual outlier markers only for visual clarity; F2 uses the requested fault-inclusive labels and no plot title.\n''')
    p=ROOT/'README.md'; t=p.read_text(); t=t.replace('(results/main/figure_2_fault_batch_rmse.png)','(figures/publishable_images/F2_fault_batch_rmse.png)')
    if '## Figure archive' not in t:
        block='''\n## Figure archive\n\nThe original result images are preserved in `results/`. Paper-ready versions are organized separately.\n\n| Figure set | Purpose |\n|---|---|\n| [images_300_dpi](figures/images_300_dpi/) | Organized copies of the previous/canonical outputs. |\n| [images_600_dpi_original_code_output](figures/images_600_dpi_original_code_output/) | 600-DPI original-output appearance plus PowerPoint. |\n| [publishable_images](figures/publishable_images/) | Publication-ready 600-DPI figures and PowerPoints; used for the README preview. |\n| [paper_figures_600dpi.py](figures/paper_figures_600dpi.py) | Reproducible Python generator. |\n| [image_manifest.csv](figures/image_manifest.csv) | Dimensions, DPI and checksums. |\n\n'''
        t=t.replace('## Supporting files\n',block+'## Supporting files\n')
    p.write_text(t)
    p=ROOT/'results/main/README.md'; t=p.read_text().replace('(figure_2_fault_batch_rmse.png)','(../../figures/publishable_images/F2_fault_batch_rmse.png)'); p.write_text(t)
    p=ROOT/'docs/FIGURES.md'; t=p.read_text();
    if 'high-resolution figure archive' not in t: t=t.replace('[Project overview](../README.md) · [Results guide](RESULTS_GUIDE.md) · [Figure provenance](figure_provenance.json)\n','[Project overview](../README.md) · [Results guide](RESULTS_GUIDE.md) · [Figure provenance](figure_provenance.json)\n\nFor manuscript preparation, see the organized [high-resolution figure archive](../figures/README.md). Original result PNGs remain unchanged in `results/`.\n')
    p.write_text(t)
    p=ROOT/'.gitattributes'; t=p.read_text();
    if '*.pptx binary' not in t: p.write_text(t.rstrip()+'\n*.pptx binary\n')

def cleanup():
    shutil.rmtree(ROOT/'paper_figures_600dpi',ignore_errors=True)
    for p in (ROOT/'IndPenSim_14_Paper_Figures_600dpi.pptx',ROOT/'IndPenSim_14_Paper_Figures_600dpi.zip'):
        if p.exists(): p.unlink()

def main():
    prep(); legacy_sets(); publication_set(); deck(D600,D600/'Original_Figure_as_Code_Output.pptx'); deck(DPUB,DPUB/'Publications_Images.pptx',True); manifest(); docs(); cleanup(); print('Built organized figure archive')
if __name__=='__main__': main()
